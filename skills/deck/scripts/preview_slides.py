"""Approximate renderer for a built deck, so every slide can be looked at without PowerPoint.

Draws each slide from the .pptx shape tree: boxes at their inch coordinates, text at its
real point size with wrapping, tables as grids, pictures from their embedded bytes. It is
a proof that nothing is blank, overlapping or running off the page, not a pixel-accurate
preview; text wrapping is estimated from average glyph width, so a line that looks clipped
here is a "check this slide" flag, not a defect on its own. Read the PNGs with the Read
tool after every build; then open the .pptx itself before hand-over.

    python3 preview_slides.py deck.pptx [--out preview/]     # writes preview/slide_N.png
"""
import argparse
import io
import textwrap
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

W, H = 13.333, 7.5


def inch(v):
    return Emu(v).inches if v is not None else 0.0


def fill_hex(sh):
    try:
        if sh.fill.type is not None and sh.fill.type == 1:
            return "#" + str(sh.fill.fore_color.rgb)
    except Exception:
        pass
    return None


def draw_text_frame(ax, sh, left, top, width):
    y = top + 0.04
    for p in sh.text_frame.paragraphs:
        if not p.text.strip():
            continue
        size = max((r.font.size.pt for r in p.runs if r.font.size), default=12)
        color = "#15180F"
        for r in p.runs:
            if r.font.color and r.font.color.type is not None:
                try:
                    color = "#" + str(r.font.color.rgb)
                except Exception:
                    pass
                break
        bold = any(r.font.bold for r in p.runs)
        cpl = max(6, int(width / (size * 0.5 / 72)))
        for line in textwrap.wrap(p.text, cpl) or [""]:
            ax.text(left, y, line, fontsize=size, color=color, va="top", ha="left",
                    fontweight="bold" if bold else "normal")
            y += size * 1.22 / 72
        y += (p.space_after.pt / 72) if p.space_after else 0.0


def draw_table(ax, sh, left, top):
    tbl = sh.table
    y = top
    for r_i, row in enumerate(tbl.rows):
        x = left
        rh = inch(row.height)
        for c_i, cell in enumerate(row.cells):
            cw = inch(tbl.columns[c_i].width)
            bg = None
            try:
                bg = "#" + str(cell.fill.fore_color.rgb)
            except Exception:
                pass
            ax.add_patch(Rectangle((x, y), cw, rh, facecolor=bg or "white",
                                   edgecolor="#D8DDC8", linewidth=0.5))
            txt = cell.text
            size = max((run.font.size.pt for p in cell.text_frame.paragraphs
                        for run in p.runs if run.font.size), default=11)
            color = "#15180F"
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    if run.font.color and run.font.color.type is not None:
                        try:
                            color = "#" + str(run.font.color.rgb)
                        except Exception:
                            pass
                    break
                break
            cpl = max(6, int((cw - 0.18) / (size * 0.5 / 72)))
            yy = y + 0.06
            for line in textwrap.wrap(txt, cpl):
                ax.text(x + 0.09, yy, line, fontsize=size, color=color, va="top")
                yy += size * 1.22 / 72
            x += cw
        y += rh


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("deck", type=Path)
    ap.add_argument("--out", type=Path, default=None, help="default: <deck dir>/preview")
    a = ap.parse_args()
    OUT = a.out or (a.deck.resolve().parent / "preview")
    OUT.mkdir(parents=True, exist_ok=True)
    prs = Presentation(str(a.deck))
    global W, H
    W, H = Emu(prs.slide_width).inches, Emu(prs.slide_height).inches
    for n, slide in enumerate(prs.slides, 1):
        fig, ax = plt.subplots(figsize=(W, H))
        ax.set_xlim(0, W)
        ax.set_ylim(H, 0)
        ax.axis("off")
        fig.subplots_adjust(0, 0, 1, 1)
        for sh in slide.shapes:
            left, top = inch(sh.left), inch(sh.top)
            width, height = inch(sh.width), inch(sh.height)
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                im = Image.open(io.BytesIO(sh.image.blob))
                ax.imshow(im, extent=(left, left + width, top + height, top))
                continue
            if getattr(sh, "has_table", False) and sh.has_table:
                draw_table(ax, sh, left, top)
                continue
            bg = fill_hex(sh)
            if bg:
                ax.add_patch(Rectangle((left, top), width, height, facecolor=bg,
                                       edgecolor="none"))
            if sh.has_text_frame and sh.text_frame.text.strip():
                draw_text_frame(ax, sh, left, top, width)
        ax.add_patch(Rectangle((0, 0), W, H, fill=False, edgecolor="#BBBBBB", lw=1))
        fig.savefig(OUT / f"slide_{n}.png", dpi=110)
        plt.close(fig)
        print(f"  {OUT / f'slide_{n}.png'}")


if __name__ == "__main__":
    main()
