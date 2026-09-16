"""Fail the build when a number reached a slide without passing through the registry.

The deck builder already enforces one direction of the skeleton's first rule: it types no
figures, it looks every one up in analysis/figures.csv by key, and an unknown key raises.
That stops a typo. It does not stop the other direction, which is the one that actually
ships wrong numbers:

  - a value hardcoded straight into a table literal, never added to the registry
  - a number written into prose because it was "obvious" at the time
  - a future edit that pastes a figure in rather than adding a row

So this scans the built .pptx and reports every numeral in slide text that does not
resolve to a registry row. It also reports the reverse: registry rows marked measured or
quoted whose value appears on no slide. Those are either a slide that lost a number or a
row that should be retired, and both are worth knowing before a workshop.

    python3 check_figures.py --deck deck.pptx --registry analysis/figures.csv [--charts charts/] [--strict]
                             [--exempt REGEX ...]

Registry columns expected: key, slide, claim, value, row_set, source, status
(status: measured | quoted | assumed | pending | blocked).

Exit 1 on any unmatched numeral, so this can gate a build. Orphaned registry rows are
reported but do not fail by default, because a chart legitimately carries its own numbers
from the CSV; pass --strict to fail on those too.

What counts as a match
----------------------
Every numeral is reduced to a canonical number: separators, spaces, currency symbols and
units stripped, then parsed as a float, so 0.810 matches 0.81 and "$3,014" matches 3,014
and "100%" matches 100. A slide numeral matches if that canonical form appears anywhere in
the registry's value, row_set or claim text. Compounds need no special handling as a
result: "3,000 to 18,000" is checked as 3000 and 18000, "15 of 20" as 15 and 20,
"0.56 MB at 2 ms and F1 0.810" as 0.56, 2 and 0.81, and each part has to be in the
registry on its own.

row_set is included because the deck legitimately quotes denominators out of it ("of the
384,271 rejected", "of the 16,292 test rows"). claim is included because it is registry-
authored text too, and labels that exist only in the registry ("Top 22", "Top 30" on a
coverage table) have no other home. That is the one deliberate relaxation here, so the report says how many
numerals leaned on it and nothing else.

This is an existence check, not a semantics check. A slide could still quote a real
registry number against the wrong claim, and 98.9% on slide 9 will match the four-country
share on slide 3 because both are 98.9. Catching that needs a human; catching a number
with no registry row at all does not.
"""
import argparse
import csv
import re
import sys
from collections import defaultdict
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


MONTHS = (r"Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec|January|February|March|"
          r"April|June|July|August|September|October|November|December")

# Spans masked out before any numeral is looked for. Each one is a place where digits
# appear on a slide without being a measurement, so there is nothing in the registry they
# could resolve to and flagging them would only train a reader to ignore this report.
EXEMPT = [
    # Cross-references to the deck's own structure: "Sections 7 through 9", "slide 15",
    # "slides 22 and 23". Navigation, not a claim about the data.
    (r"\bslides?\s+\d+(?:\s*(?:and|to|through|,|-|&)\s*\d+)*", "slide reference"),
    (r"\bsections?\s+\d+(?:\s*(?:and|to|through|,|-|&)\s*\d+)*", "section reference"),
    # Work-plan references. "sub-task 2" names an owner, not a quantity.
    (r"\bsub-?tasks?\s+\d+(?:\s*(?:and|,|-)\s*\d+)*", "sub-task reference"),
    # Checkpoint identifiers. M0, M1 and M2 are the names of the three pipeline stages,
    # fixed in the skeleton so every slide can refer to them the same way.
    (r"\bM[012]\b", "checkpoint identifier"),
    # Dates and years. "Aug 2024 to Jul 2026", "9 August", "August 2026", "31 Aug 2026".
    # The period covered is a registry row; the individual dates inside prose are not.
    (r"\b(?:19|20)\d{2}\b", "year"),
    (rf"\b\d{{1,2}}\s+(?:{MONTHS})\b", "day and month"),
    (rf"\b(?:{MONTHS})\s+\d{{1,2}}\b", "month and day"),
    # Model names carrying digits. Version numbers are identity, not measurement, and
    # "Sonnet 4.6" would otherwise report a 4.6 that can never have a registry row.
    # Both the full name and the short form the callouts use: Qwen3-VL-235B-A22B, Qwen-235B.
    # Generic: a model family name followed by a version or size token. Add project-specific
    # names with --exempt on the command line rather than editing this list.
    (r"Qwen\d?(?:[\s-]*VL)?(?:[\s-]*\d+B)?(?:-A\d+B)?", "model name"),
    (r"(?:Sonnet|Opus|Haiku|Claude)[\s-]*\d(?:\.\d)?", "model name"),
    (r"Gemini[\s-]*\d(?:\.\d)?(?:\s*(?:Flash|Pro))?", "model name"),
    (r"Gemma[\s-]*\d(?:-?e?\d+B)?", "model name"),
    (r"Kimi[\s-]*K\d(?:\.\d)?", "model name"),
    (r"MobileNetV\d(?:-small)?", "model name"),
    (r"DaViT-(?:Small|Base|Tiny)", "model name"),
    (r"GPT-?[45][a-z]?(?:[\s-]*mini)?\b", "model name"),
    (r"Pixtral(?:\s+Large)?", "model name"),
    (r"Llama[\s-]*\d+(?:\.\d+)?", "model name"),
    (r"(?:Indic)?Whisper(?:[\s-]*(?:v\d|large|medium|small))?", "model name"),
    (r"(?:Sarvam|Conformer|Canary|Parakeet|wav2vec|Deepgram)[\s-]*[\w.]*\d[\w.]*", "model name"),
    (r"(?:EfficientNet|ConvNeXt)[\s-]*[\w.-]*\d[\w.-]*", "model name"),
    # Slide furniture: the "3 / 8" page footer, "Slide 3 of 8", and the zero-padded labels
    # ("01", "02") the numbered-point block draws in front of each point.
    (r"\b\d{1,2}\s*/\s*\d{1,2}\b", "page footer"),
    (r"\bslide\s+\d{1,2}\s+of\s+\d{1,2}\b", "page footer"),
    (r"(?<![\d.,])0\d(?![\d.,%])", "point label"),
    # Ordinals in prose: "50th percentile", "the twentieth class". The percentile itself is
    # the registry row; the ordinal naming it is not a second number.
    (r"\b\d+(?:st|nd|rd|th)\b", "ordinal"),
    # Metric and percentile labels whose name contains a digit: F1, p50, p95.
    (r"\bF1\b", "metric name"),
    (r"\bp\d{1,3}\b", "percentile label"),
]
EXEMPT = [(re.compile(p, re.IGNORECASE), why) for p, why in EXEMPT]

# A numeral plus the unit that belongs to it. Anchored so it cannot start inside a word
# or a version string, and cannot stop halfway through "1,163,658" or "0.810".
NUM = re.compile(
    r"(?<![A-Za-z0-9.])"
    r"(\$?\d(?:[\d,]*\d)?(?:\.\d+)?)"      # cannot end on a separator, so "42," is "42"
    r"(\s*(?:%|ms|MB|KB|GB|¢))?"
    r"(?!\.?\d)"
)


def canon(raw):
    """One number, reduced so that presentation differences stop mattering. Returns None
    when the token is not really a number (a bare comma-mangled id, say)."""
    s = raw.strip().lstrip("$").replace(",", "").replace(" ", "")
    for unit in ("%", "ms", "MB", "KB", "GB", "¢"):
        if s.endswith(unit):
            s = s[: -len(unit)]
    try:
        return f"{float(s):.6f}".rstrip("0").rstrip(".") or "0"
    except ValueError:
        return None


def mask(text):
    """The text with exempt spans blanked to the same length, so offsets into the original
    still line up and the reported context reads as it does on the slide."""
    out = list(text)
    for pat, _ in EXEMPT:
        for m in pat.finditer(text):
            for i in range(m.start(), m.end()):
                out[i] = " "
    return "".join(out)


def numerals(text):
    """(raw token including unit, canonical form, start, end) for each numeral."""
    found = []
    for m in NUM.finditer(text):
        c = canon(m.group(1))
        if c is not None:
            found.append((m.group(0).strip(), c, m.start(), m.end()))
    return found


def read_registry(path):
    with path.open(newline="") as fh:
        rows = list(csv.DictReader(fh))
    if not rows:
        sys.exit(f"{path} has no rows")
    return rows


def registry_index(rows):
    """Canonical numeral -> set of the registry fields it was found in."""
    idx = defaultdict(set)
    for r in rows:
        for field in ("value", "row_set", "claim"):
            for _, c, _, _ in numerals(r.get(field) or ""):
                idx[c].add(field)
    return idx


def slide_texts(deck):
    """[(slide number, [text block, ...], has_picture)] over every shape, table row and
    grouped shape. A table contributes one block per row, cells joined, because a bare
    "83.4%" cell reported on its own tells a reader nothing about which row it came from,
    and the whole row is what lets them judge it in one read."""
    prs = Presentation(str(deck))
    out = []
    for n, slide in enumerate(prs.slides, 1):
        blocks, pic = [], False
        def walk(shapes):
            nonlocal pic
            for sh in shapes:
                if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                    walk(sh.shapes)
                    continue
                if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    pic = True
                if getattr(sh, "has_table", False) and sh.has_table:
                    for row in sh.table.rows:
                        cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                        if any(cells):
                            blocks.append(" | ".join(cells))
                    continue
                if sh.has_text_frame and sh.text_frame.text.strip():
                    blocks.append(sh.text_frame.text.strip())
        walk(slide.shapes)
        out.append((n, blocks, pic))
    return out


def context(block, start, end, width=88):
    """The token in enough of its sentence for a human to judge it in one read."""
    flat = block.replace("\n", " ").replace(" ", " ")
    pad = max(0, (width - (end - start)) // 2)
    a, b = max(0, start - pad), min(len(flat), end + pad)
    return ("…" if a else "") + " ".join(flat[a:b].split()) + ("…" if b < len(flat) else "")


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--deck", type=Path, required=True)
    ap.add_argument("--registry", type=Path, required=True)
    ap.add_argument("--charts", type=Path, default=None,
                    help="folder holding gen_*.py and their .dot outputs (default: <deck dir>/charts)")
    ap.add_argument("--exempt", action="append", default=[], metavar="REGEX",
                    help="extra span to mask before scanning (repeatable), e.g. a project model name")
    ap.add_argument("--strict", action="store_true",
                    help="also exit non-zero on registry rows no slide quotes")
    args = ap.parse_args()
    for pat in args.exempt:
        EXEMPT.append((re.compile(pat, re.IGNORECASE), "command-line exempt"))

    rows = read_registry(args.registry)
    idx = registry_index(rows)
    slides = slide_texts(args.deck)

    unmatched = []          # (slide, token, context)
    claim_only = []         # matched, but only because claim text was allowed
    on_slides = set()       # every canonical numeral anywhere in slide text
    seen = set()

    for n, blocks, _ in slides:
        for block in blocks:
            masked = mask(block)
            for raw, c, start, end in numerals(block):
                on_slides.add(c)
            for raw, c, start, end in numerals(masked):
                where = idx.get(c)
                if not where:
                    key = (n, raw, context(block, start, end))
                    if key not in seen:
                        seen.add(key)
                        unmatched.append(key)
                elif where == {"claim"}:
                    claim_only.append((n, raw))

    # The reverse direction. A row whose value carries numerals is present if every one of
    # them shows up in slide text; a row whose value is words ("Aug 2024 to Jul 2026",
    # "pending") is present if that text appears on a slide. Only measured and quoted rows
    # are checked: pending and blocked rows are meant to render as the word "pending".
    all_text = " ¶ ".join(b for _, blocks, _ in slides for b in blocks).lower()
    pictured = {n for n, _, pic in slides if pic}

    # Numerals that reach the reader through a generated diagram rather than slide text.
    #
    # The .dot files in charts/ are build artifacts: charts/gen_*.py rewrite them from the
    # same CSVs the registry rows cite, so a number in one re-renders when the data
    # changes. That is not a number "baked into an image" in the sense the skeleton
    # forbids, which is a figure nobody can update without redrawing. A hand-authored .dot
    # would be exactly that, so these are reported as their own category rather than
    # silently accepted, and only .dot files a generator writes are trusted.
    chart_dir = args.charts or (args.deck.parent / "charts")
    generated = set()
    for gen in sorted(chart_dir.glob("gen_*.py")):
        src = gen.read_text()
        generated.update(re.findall(r'render\(\s*["\']([\w-]+)["\']', src))
        generated.update(re.findall(r'OUT\s*=\s*HERE\s*/\s*["\']([\w-]+)["\']', src))
    in_charts = set()
    for name in sorted(generated):
        dot = chart_dir / f"{name}.dot"
        if dot.exists():
            # DOT label text carries literal \n escapes, and a digit sitting straight
            # after the "n" has no word boundary in front of it, so the numeral regex
            # skips it. Turn the escapes into spaces before scanning.
            txt = dot.read_text().replace("\\n", " ").replace("\\l", " ")
            for _, c, _, _ in numerals(txt):
                in_charts.add(c)

    orphans = []
    charted = []
    for r in rows:
        if r.get("status") not in ("measured", "quoted"):
            continue
        value = (r.get("value") or "").strip()
        nums = numerals(value)
        if nums:
            missing = [c for _, c, _, _ in nums if c not in on_slides]
            present = not missing
        else:
            present = value.lower() in all_text
        if not present:
            try:
                sl = int(r.get("slide") or 0)
            except ValueError:
                sl = 0
            if nums and all(c in in_charts for _, c, _, _ in nums):
                charted.append((sl, r["key"], value))
            else:
                orphans.append((sl, r["key"], value, sl in pictured))

    # ------------------------------------------------------------------------ report
    print(f"deck      {args.deck}")
    print(f"registry  {args.registry}  ({len(rows)} rows, "
          f"{sum(1 for r in rows if r.get('status') in ('measured', 'quoted'))} "
          f"measured or quoted)")
    print(f"slides    {len(slides)}")
    print()

    if unmatched:
        print(f"UNMATCHED NUMERALS ({len(unmatched)}) : on a slide, not in the registry")
        print()
        for n, raw, ctx in sorted(unmatched, key=lambda t: (t[0], t[1])):
            print(f'  slide {n}: "{raw}" in "{ctx}"')
        print()
    else:
        print("UNMATCHED NUMERALS (0) : every numeral on every slide resolves to a row")
        print()

    if orphans:
        print(f"REGISTRY ROWS NO SLIDE QUOTES ({len(orphans)}) : a slide lost the number, "
              f"or the row should be retired")
        print()
        for sl, key, value, pic in sorted(orphans):
            note = "  (that slide carries an image, so a chart may hold it)" if pic else ""
            print(f"  slide {sl}: {key} = {value}{note}")
        print()

    if charted:
        print(f"CARRIED BY A GENERATED DIAGRAM ({len(charted)}) : not in slide text, but "
              f"every numeral is in a .dot a generator rewrites from the cited CSV, so it "
              f"re-renders with the data")
        print()
        for sl, key, value in sorted(charted):
            print(f"  slide {sl}: {key} = {value}")
        print()

    if claim_only:
        counts = defaultdict(int)
        for n, raw in claim_only:
            counts[(n, raw)] += 1
        print(f"MATCHED ONLY VIA A REGISTRY claim FIELD ({len(counts)}) : allowed, but "
              f"these have no value or row_set behind them")
        for (n, raw), _ in sorted(counts.items()):
            print(f"  slide {n}: {raw}")
        print()

    failed = bool(unmatched) or (args.strict and bool(orphans))
    print("FAIL" if failed else "PASS")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
