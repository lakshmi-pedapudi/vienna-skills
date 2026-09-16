"""Word gate for a built deck: em-dashes, AI-tell phrases, board-deck jargon, banned slide
words, internal section codes, and personal names (from an optional names file).

Scans slide text, table cells and speaker notes. Exit 1 on any hit.

    python3 check_words.py deck.pptx [--names names.txt] [--allow REGEX ...] [--no-notes]

names.txt: one name per line (team members who must not appear on a slide face). The
presenter's name on the cover is the usual --allow.
"""
import argparse
import re
import sys
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

RULES = [
    ("em-dash", r"—|&mdash;|&#8212;"),
    ("AI tell", r"\b(?:dive into|delve|leverag\w*|robust|seamless\w*|unlock\w*|genuinely|paramount|"
                r"it is worth noting|in today.s|underscor\w* the need|sweet spot|comprehensive)\b"),
    ("jargon", r"\b(?:blast radius|mitigation in flight|force multiplier|fan-?out|cohort|"
               r"surface area|operationali[sz]e\w*|headline)\b"),
    ("banned slide word", r"\bTheme\s*\d*\b|\bWhy (?:it|this) matters\b|\bWhat this means\b|"
                          r"\bRisk surface\b|\bV2 product pivot\b|\bCross-link:"),
    ("section code", r"\b(?:[A-E]\d(?:\.\w+)?|F\d\.\w+|T\d\.\d|D\.3\.\w+)\b"),
    ("invented metric word", r"\b(?:under the rule|verdict|admission|axes)\b"),
]


def title_of(slide):
    """Topmost text shape on a blank-layout slide is the title (or the eyebrow above it)."""
    best = None
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip() and sh.top is not None:
            if best is None or sh.top < best.top:
                best = sh
    if best is None:
        return ""
    paras = [p.text.strip() for p in best.text_frame.paragraphs if p.text.strip()]
    return paras[-1] if paras else ""   # eyebrow sits above the title in the same box


def title_problems(title):
    words = title.split()
    if len(words) > 4:
        yield f"title has {len(words)} words (1 to 4 plain nouns expected)"
    if title.endswith((".", "?", "!")):
        yield "title ends with sentence punctuation"
    if re.match(r"(?i)^(?:what|how|why|when|where|two ways|the )", title):
        yield "title reads as a sentence or narrative label"
    if re.search(r"\b(?:not|vs\.?|versus)\b", title, re.I):
        yield "'X not Y' or versus construction in a title"


def texts(prs, with_notes):
    for n, slide in enumerate(prs.slides, 1):
        def walk(shapes, where):
            for sh in shapes:
                if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                    yield from walk(sh.shapes, where)
                    continue
                if getattr(sh, "has_table", False) and sh.has_table:
                    for row in sh.table.rows:
                        for c in row.cells:
                            if c.text.strip():
                                yield n, where, c.text
                    continue
                if sh.has_text_frame and sh.text_frame.text.strip():
                    yield n, where, sh.text_frame.text
        yield from walk(slide.shapes, "slide")
        if with_notes and slide.has_notes_slide:
            t = slide.notes_slide.notes_text_frame.text
            if t.strip():
                yield n, "notes", t


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("deck")
    ap.add_argument("--names", help="file with one personal name per line")
    ap.add_argument("--allow", action="append", default=[], metavar="REGEX")
    ap.add_argument("--no-notes", action="store_true", help="skip speaker notes")
    a = ap.parse_args()

    rules = [(k, re.compile(p, re.IGNORECASE | re.MULTILINE)) for k, p in RULES]
    if a.names:
        names = [l.strip() for l in open(a.names) if l.strip()]
        if names:
            rules.append(("personal name", re.compile(r"\b(?:" + "|".join(map(re.escape, names)) + r")\b")))
    allow = [re.compile(p, re.IGNORECASE) for p in a.allow]

    hits = 0
    prs = Presentation(a.deck)
    for n, slide in enumerate(prs.slides, 1):
        t = title_of(slide)
        for prob in title_problems(t):
            hits += 1
            print(f"slide {n} (title) [{prob}]: {t[:90]!r}")
    for n, where, t in texts(prs, not a.no_notes):
        for kind, rx in rules:
            if kind == "section code" and where == "notes":
                continue  # codes are allowed in notes as pointers to working docs
            for m in rx.finditer(t):
                frag = " ".join(t[max(0, m.start() - 40):m.end() + 40].split())
                if any(ax.search(frag) for ax in allow):
                    continue
                hits += 1
                print(f"slide {n} ({where}) [{kind}] {m.group(0)!r}: ...{frag}...")
    print(f"\n{hits} hit(s)")
    sys.exit(1 if hits else 0)


if __name__ == "__main__":
    main()
