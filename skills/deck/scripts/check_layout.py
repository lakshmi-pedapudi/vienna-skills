"""Flag overlapping shapes and overfull text in the built deck.

There is no PowerPoint renderer on this machine, so the deck cannot be looked at here.
This checks the two things that go wrong without a renderer to catch them: shapes whose
declared boxes overlap, and text that will not fit the box it was given.

Both are approximations of what PowerPoint will do, and the text estimate is the rougher
of the two, so treat a warning as "open this slide and look" rather than as a defect.
Tables are the main risk: row_h is a minimum, and a cell whose text wraps to two lines
makes the whole table taller than it was declared, which is how a table ends up sitting
on top of a footnote.

    python3 check_layout.py deck.pptx [--foot-rule 6.48]

--foot-rule is the y position (inches) of the rule above the footnote; anything crossing it
collides with the footnote. Default matches the house primitives (FOOT_T 6.62 minus 0.14).
"""
import argparse
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

# The rule above the footnote. Anything crossing this line collides with the footnote.
FOOT_RULE = 6.48


def inches(v):
    return Emu(v).inches if v is not None else None


def rect(sh):
    try:
        return (inches(sh.left), inches(sh.top),
                inches(sh.left) + inches(sh.width),
                inches(sh.top) + inches(sh.height))
    except TypeError:
        return None


def text_of(sh):
    if not sh.has_text_frame:
        return ""
    return "\n".join(p.text for p in sh.text_frame.paragraphs)


def est_height(sh):
    """Rough rendered height of a text frame, in inches.

    Characters per line comes from the average glyph width of Arial, about 0.5 em, and
    line height from 1.22 times the point size, which is the line_spacing this deck sets.
    """
    if not sh.has_text_frame:
        return None
    w = inches(sh.width)
    total = 0.0
    for p in sh.text_frame.paragraphs:
        txt = p.text
        if not txt:
            continue
        size = max((r.font.size.pt for r in p.runs if r.font.size), default=12)
        cpl = max(8, int(w / (size * 0.5 / 72)))
        lines = max(1, -(-len(txt) // cpl))
        total += lines * (size * 1.22 / 72)
        if p.space_after:
            total += p.space_after.pt / 72
    return total


def table_min_height(sh):
    """Declared table height, plus the extra a wrapping cell will force.

    A cell whose text is wider than its column wraps, and PowerPoint grows the row to
    fit. That growth is what pushes tables into footnotes, so it is estimated here
    rather than ignored.
    """
    tbl = sh.table
    extra = 0.0
    for r_i, row in enumerate(tbl.rows):
        worst = 1
        for c_i, cell in enumerate(row.cells):
            cw = inches(tbl.columns[c_i].width) - 0.18   # cell margins
            txt = cell.text_frame.text
            size = max((run.font.size.pt for p in cell.text_frame.paragraphs
                        for run in p.runs if run.font.size), default=12)
            cpl = max(6, int(cw / (size * 0.5 / 72)))
            worst = max(worst, -(-len(txt) // cpl))
        if worst > 1:
            rh = inches(row.height)
            need = worst * (12 * 1.22 / 72) + 0.06
            extra += max(0.0, need - rh)
    return inches(sh.height) + extra


def main():
    global FOOT_RULE
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("deck", type=Path)
    ap.add_argument("--foot-rule", type=float, default=FOOT_RULE)
    args = ap.parse_args()
    FOOT_RULE = args.foot_rule
    prs = Presentation(str(args.deck))
    problems = 0
    for i, slide in enumerate(prs.slides, 1):
        notes = []
        boxes = []
        for sh in slide.shapes:
            r = rect(sh)
            if r is None:
                continue
            bottom = r[3]
            label = sh.shape_type
            if sh.has_table:
                grown = table_min_height(sh)
                bottom = r[1] + grown
                if grown > inches(sh.height) + 0.02:
                    notes.append(f"table at y={r[1]:.2f} grows "
                                 f"{inches(sh.height):.2f} -> {grown:.2f} in")
            elif sh.has_text_frame and sh.text_frame.text.strip():
                eh = est_height(sh)
                if eh and eh > inches(sh.height) + 0.12:
                    notes.append(f"text at y={r[1]:.2f} needs ~{eh:.2f} in of "
                                 f"{inches(sh.height):.2f}: "
                                 f"{text_of(sh)[:52]!r}")
            boxes.append((r[0], r[1], r[2], bottom, label, sh))
            if bottom > FOOT_RULE + 0.02 and sh.has_table:
                notes.append(f"table bottom {bottom:.2f} crosses the footnote rule "
                             f"at {FOOT_RULE}")
            if bottom > FOOT_RULE + 0.02 and sh.shape_type == 13:
                notes.append(f"picture bottom {bottom:.2f} crosses the footnote rule")

        # Overlap between content boxes. Background fills, rules and the tint behind a
        # callout are meant to sit under things, so only flag pairs that both carry
        # content.
        content = [b for b in boxes
                   if (b[5].has_table or b[5].shape_type == 13
                       or (b[5].has_text_frame and b[5].text_frame.text.strip()))]
        for a in range(len(content)):
            for b in range(a + 1, len(content)):
                x1, y1, x2, y2, _, sa = content[a]
                u1, v1, u2, v2, _, sb = content[b]
                ox = min(x2, u2) - max(x1, u1)
                oy = min(y2, v2) - max(y1, v1)
                if ox > 0.25 and oy > 0.25:
                    ta = (text_of(sa)[:26] or "table/picture").replace("\n", " ")
                    tb = (text_of(sb)[:26] or "table/picture").replace("\n", " ")
                    notes.append(f"overlap {oy:.2f} in tall: {ta!r} and {tb!r}")
        if notes:
            problems += len(notes)
            print(f"slide {i}")
            for n in notes:
                print(f"    {n}")
    print(f"\n{problems} things to look at")


if __name__ == "__main__":
    main()
